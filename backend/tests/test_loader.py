from pathlib import Path

import pytest
from docx import Document
from docx.shared import Inches

from eval.synthetic_gen import _image, _run, _table, _tracked
from rug import loaders


def _save(doc, path: Path) -> Path:
    doc.save(str(path))
    return path


def test_headings_build_path_and_title(tmp_path):
    d = Document()
    d.add_paragraph("My SOW", style="Title")
    d.add_heading("1 Scope", level=1)
    d.add_heading("1.1 In scope", level=2)
    d.add_paragraph("Build the thing.")
    d.add_heading("2 Fees", level=1)
    d.add_paragraph("Ten dollars.")
    doc = loaders.load(_save(d, tmp_path / "a.docx"))

    assert (doc.title, doc.title_source) == ("My SOW", "style")
    assert [(s.heading_path, s.text) for s in doc.sections] == [
        (("1 Scope", "1.1 In scope"), "Build the thing."),
        (("2 Fees",), "Ten dollars."),
    ]


def test_core_title_wins_and_filename_fallback(tmp_path):
    d = Document()
    d.core_properties.title = "Core Title"
    d.add_paragraph("Style Title", style="Title")
    assert loaders.load(_save(d, tmp_path / "a.docx")).title == "Core Title"

    d2 = Document()
    d2.add_paragraph("just text")
    doc = loaders.load(_save(d2, tmp_path / "b.docx"))
    assert (doc.title, doc.title_source) == (None, "filename")


def test_table_rows_keep_headers(tmp_path):
    d = Document()
    _table(d, ["Milestone", "Date"], [["UAT", "1 May"], ["Go-live", "2 June"]])
    [section] = loaders.load(_save(d, tmp_path / "t.docx")).sections
    assert section.kind == "table"
    assert section.text.splitlines() == [
        "Milestone | Date",
        "Milestone: UAT; Date: 1 May",
        "Milestone: Go-live; Date: 2 June",
    ]


def test_merged_cells_not_duplicated(tmp_path):
    d = Document()
    t = d.add_table(rows=2, cols=3)
    t.rows[0].cells[0].text, t.rows[0].cells[2].text = "A", "C"
    t.rows[0].cells[0].merge(t.rows[0].cells[1])  # horizontal merge -> one gridSpan cell
    t.rows[1].cells[0].text = "x"
    [section] = loaders.load(_save(d, tmp_path / "m.docx")).sections
    header = section.text.splitlines()[0]
    assert header.count("A") == 1


def test_tracked_changes_read_as_accepted(tmp_path):
    d = Document()
    p = d.add_paragraph("Pay within ")
    p._p.append(_tracked("w:del", "30 days", 1))
    p._p.append(_tracked("w:ins", "45 days", 2))
    p._p.append(_run("."))
    [section] = loaders.load(_save(d, tmp_path / "tc.docx")).sections
    assert section.text == "Pay within 45 days."


def test_ocr_image_text_and_small_images_skipped(tmp_path):
    d = Document()
    d.add_heading("Architecture", level=1)
    d.add_picture(_image("Primary site:\nMelbourne DC2", (900, 300)), width=Inches(5))
    d.add_picture(_image("LOGO", (60, 60)), width=Inches(0.5))
    doc = loaders.load(_save(d, tmp_path / "img.docx"))
    image_sections = [s for s in doc.sections if s.kind == "image_text"]
    assert len(image_sections) == 1  # the 60x60 logo is below ocr_min_px
    assert "Melbourne DC2" in image_sections[0].text
    assert image_sections[0].heading_path == ("Architecture",)


def test_not_a_zip_raises(tmp_path):
    p = tmp_path / "fake.docx"
    p.write_text("hello")
    with pytest.raises(ValueError):
        loaders.load(p)


def test_registry_is_upload_allow_list():
    assert loaders.supported_extensions() == frozenset({".docx"})
    assert loaders.get_loader("X.DOCX").looks_valid(b"PK\x03\x04rest")
    assert not loaders.get_loader("x.docx").looks_valid(b"%PDF-1.7")
    with pytest.raises(loaders.UnsupportedFormat):
        loaders.get_loader("report.pdf")
