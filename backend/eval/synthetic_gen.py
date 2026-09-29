"""Generate a synthetic .docx corpus that exercises the hard cases.

All companies, IDs and figures are fictional. `FACTS` lists, per file, strings that must be
findable after indexing (and `ABSENT` strings that must not be), so tests and the eval
runner can check extraction without an LLM.
"""

import io
import os
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.oxml.xmlchemy import BaseOxmlElement
from docx.shared import Inches
from PIL import Image, ImageDraw, ImageFont

# Relative path -> strings that must appear in the indexed chunk text.
FACTS: dict[str, list[str]] = {
    "sales/Boost Connect SR-1098 SOW v1.docx": [
        "network assessment of three sites",
        "AUD 38,000",
    ],
    "sales/Boost Connect SR-1098 SOW v2 FINAL.docx": [
        "Wi-Fi 7 rollout across five sites",
        "24/7 network operations centre monitoring for 12 months",
        "AUD 184,500",
        "30 June 2027",
        "10 business days",
        "monthly in arrears",
    ],
    "sales/Boost Connect SR-1098 CR-01.docx": ["two additional sites", "AUD 18,400", "six weeks"],
    "sales/Boost Connectivity SR-1089 SOW.docx": [
        "SD-WAN migration",
        "AUD 96,000",
        "Meridian Networks",
    ],
    "delivery/Harbor Logistics SR-2210 SOW.docx": [
        "UAT sign-off",
        "19 February 2027",
        "Stockline",
        "4 weeks",
    ],
    "delivery/Pinecrest Health SR-3301 SOW.docx": [
        "Primary datacentre",
        "Melbourne DC2",
        "15 minutes",
        "three waves",
    ],
    "delivery/Atlas Retail SR-4410 SOW.docx": ["net 45 days", "five stores per week"],
    "delivery/Northwind Foods SR-5120 SOW.docx": ["cold-chain sensor", "five years", "Dandenong"],
    "legal/Boost Connect MSA 2025.docx": [
        "liability is capped at AUD 2,000,000",
        "laws of Victoria",
        "60 days written notice",
    ],
    "legal/Contoso Mutual NDA.docx": ["three (3) years", "New South Wales"],
}
# Strings that must NOT be indexed (tracked deletions).
ABSENT: dict[str, list[str]] = {
    "delivery/Atlas Retail SR-4410 SOW.docx": ["net 30 days"],
}
TITLES: dict[str, str] = {
    "delivery/Northwind Foods SR-5120 SOW.docx": "Northwind Foods Cold-Chain Monitoring SOW",
}


def _image(text: str, size: tuple[int, int]) -> io.BytesIO:
    img = Image.new("RGB", size, "white")
    draw = ImageDraw.Draw(img)
    font = ImageFont.load_default(size=max(12, size[1] // 8))
    draw.rectangle([4, 4, size[0] - 5, size[1] - 5], outline="black", width=3)
    draw.multiline_text((24, 24), text, fill="black", font=font, spacing=12)
    buf = io.BytesIO()
    img.save(buf, "PNG")
    buf.seek(0)
    return buf


def _run(text: str, deleted: bool = False) -> BaseOxmlElement:
    r = OxmlElement("w:r")
    t = OxmlElement("w:delText" if deleted else "w:t")
    t.set(qn("xml:space"), "preserve")
    t.text = text
    r.append(t)
    return r


def _tracked(tag: str, text: str, rid: int) -> BaseOxmlElement:
    el = OxmlElement(tag)
    el.set(qn("w:id"), str(rid))
    el.set(qn("w:author"), "Reviewer")
    el.set(qn("w:date"), "2026-01-15T10:00:00Z")
    el.append(_run(text, deleted=tag == "w:del"))
    return el


def _table(doc, header: list[str], rows: list[list[str]]) -> None:
    t = doc.add_table(rows=1, cols=len(header))
    for cell, h in zip(t.rows[0].cells, header, strict=True):
        cell.text = h
    for row in rows:
        for cell, v in zip(t.add_row().cells, row, strict=True):
            cell.text = v


def _sow(title: str, client: str, overview: str, scope: list[str]):
    doc = Document()
    doc.add_paragraph(title, style="Title")
    doc.add_heading("1 Overview", level=1)
    doc.add_paragraph(overview)
    doc.add_paragraph(
        f"This Statement of Work is entered into between Example Integrators and {client}."
    )
    doc.add_heading("2 Scope of Work", level=1)
    doc.add_heading("2.1 In scope", level=2)
    for item in scope:
        doc.add_paragraph(item, style="List Bullet")
    doc.add_heading("2.2 Out of scope", level=2)
    doc.add_paragraph("Hardware procurement and end-user device support are out of scope.")
    return doc


def build(out: Path) -> list[Path]:
    """Write the corpus under `out` and return the created paths."""
    docs: dict[str, object] = {}

    d = _sow(
        "Boost Connect SR-1098 Statement of Work",
        "Boost Connect Pty Ltd",
        "Draft scope: a network assessment of three sites to inform a later Wi-Fi design.",
        ["Site survey at three sites", "Assessment report with recommendations"],
    )
    d.add_heading("3 Fees", level=1)
    d.add_paragraph("The assessment is offered at a fixed fee of AUD 38,000.")
    docs["sales/Boost Connect SR-1098 SOW v1.docx"] = d

    d = _sow(
        "Boost Connect SR-1098 Statement of Work",
        "Boost Connect Pty Ltd",
        "The engagement delivers a Wi-Fi 7 rollout across five sites and "
        "24/7 network operations centre monitoring for 12 months.",
        [
            "Detailed Wi-Fi 7 design for five sites",
            "Installation and configuration of access points",
            "24/7 NOC monitoring and incident response",
        ],
    )
    d.add_heading("3 Fees", level=1)
    _table(
        d,
        ["Item", "Amount"],
        [
            ["Design and rollout", "AUD 142,000"],
            ["NOC monitoring (12 months)", "AUD 42,500"],
            ["Total", "AUD 184,500"],
        ],
    )
    d.add_heading("4 Schedule", level=1)
    d.add_paragraph("Go-live is targeted for 30 June 2027.")
    d.add_heading("5 Acceptance", level=1)
    d.add_paragraph("Acceptance testing takes 10 business days after each site is completed.")
    d.add_heading("6 Payment", level=1)
    d.add_paragraph("Invoices are issued monthly in arrears.")
    docs["sales/Boost Connect SR-1098 SOW v2 FINAL.docx"] = d

    d = Document()
    d.add_paragraph("Change Request CR-01 to SR-1098", style="Title")
    d.add_heading("Change description", level=1)
    d.add_paragraph("Extend the Wi-Fi 7 rollout to two additional sites (Geelong and Ballarat).")
    d.add_heading("Commercial impact", level=1)
    d.add_paragraph("The change adds AUD 18,400 to the SR-1098 contract value.")
    d.add_paragraph("It also adds six weeks to the delivery schedule.")
    docs["sales/Boost Connect SR-1098 CR-01.docx"] = d

    d = _sow(
        "Boost Connectivity SR-1089 Statement of Work",
        "Boost Connectivity Ltd",
        "An SD-WAN migration for twelve branch offices.",
        ["SD-WAN design", "Branch cut-over", "Decommission MPLS circuits"],
    )
    d.add_heading("3 Commercials", level=1)
    d.add_paragraph(
        "The fixed fee is AUD 96,000. The SD-WAN appliances are supplied by Meridian Networks."
    )
    docs["sales/Boost Connectivity SR-1089 SOW.docx"] = d

    # Fact only in a table.
    d = _sow(
        "Harbor Logistics SR-2210 Statement of Work",
        "Harbor Logistics",
        "Warehouse management system integration.",
        ["WMS integration", "Data migration"],
    )
    d.add_heading("3 Milestones", level=1)
    _table(
        d,
        ["Milestone", "Date", "Owner"],
        [
            ["Design complete", "5 December 2026", "Example Integrators"],
            ["UAT sign-off", "19 February 2027", "Harbor Logistics"],
            ["Go-live", "8 March 2027", "Joint"],
        ],
    )
    d.add_heading("4 Data and support", level=1)
    d.add_paragraph("Data is migrated from the legacy Stockline system.")
    d.add_paragraph("Hypercare lasts 4 weeks after go-live.")
    docs["delivery/Harbor Logistics SR-2210 SOW.docx"] = d

    # Fact only in an image; plus a tiny logo that must be skipped.
    d = _sow(
        "Pinecrest Health SR-3301 Statement of Work",
        "Pinecrest Health",
        "Hosting platform refresh. The target architecture is shown in the diagram below.",
        ["Platform build", "Migration of clinical systems"],
    )
    d.add_heading("3 Target architecture", level=1)
    d.add_picture(
        _image("Primary datacentre:\nMelbourne DC2\nDR site: Sydney", (900, 400)), width=Inches(6)
    )
    d.add_picture(_image("X", (40, 40)), width=Inches(0.4))
    d.add_heading("4 Recovery and migration", level=1)
    d.add_paragraph("The recovery point objective (RPO) is 15 minutes.")
    d.add_paragraph("Clinical systems are migrated in three waves.")
    docs["delivery/Pinecrest Health SR-3301 SOW.docx"] = d

    # Tracked changes: insertion accepted, deletion dropped.
    d = _sow(
        "Atlas Retail SR-4410 Statement of Work",
        "Atlas Retail",
        "Point-of-sale network upgrade for 40 stores.",
        ["Store network upgrade", "POS VLAN segmentation"],
    )
    d.add_heading("3 Commercial terms", level=1)
    p = d.add_paragraph("Invoices are payable within ")
    p._p.append(_tracked("w:del", "net 30 days", 1))
    p._p.append(_tracked("w:ins", "net 45 days", 2))
    p._p.append(_run(" of the invoice date."))
    d.add_heading("4 Rollout", level=1)
    d.add_paragraph("The store rollout proceeds at five stores per week.")
    docs["delivery/Atlas Retail SR-4410 SOW.docx"] = d

    d = _sow(
        "Northwind Foods SR-5120 SOW",
        "Northwind Foods",
        "Deployment of cold-chain sensor monitoring in two distribution centres.",
        ["Sensor installation", "Dashboard configuration"],
    )
    d.add_heading("3 Sites and equipment", level=1)
    d.add_paragraph("The two distribution centres are located in Dandenong and Truganina.")
    d.add_paragraph("The sensors have a battery life of five years.")
    docs["delivery/Northwind Foods SR-5120 SOW.docx"] = d

    d = Document()
    d.add_paragraph("Master Services Agreement", style="Title")
    d.add_heading("12 Limitation of liability", level=1)
    d.add_paragraph(
        "Each party's aggregate liability is capped at AUD 2,000,000 per contract year."
    )
    d.add_heading("13 Governing law", level=1)
    d.add_paragraph("This agreement is governed by the laws of Victoria.")
    d.add_heading("14 Termination", level=1)
    d.add_paragraph("Either party may terminate for convenience on 60 days written notice.")
    docs["legal/Boost Connect MSA 2025.docx"] = d

    d = Document()
    d.add_paragraph("Mutual Non-Disclosure Agreement", style="Title")
    d.add_heading("Term", level=1)
    d.add_paragraph("Confidentiality obligations survive for three (3) years after disclosure.")
    d.add_heading("Governing law", level=1)
    d.add_paragraph("This agreement is governed by the laws of New South Wales.")
    docs["legal/Contoso Mutual NDA.docx"] = d

    written: list[Path] = []
    for i, (rel, doc) in enumerate(docs.items()):
        if rel in TITLES:
            doc.core_properties.title = TITLES[rel]  # type: ignore[attr-defined]
        path = out / rel
        path.parent.mkdir(parents=True, exist_ok=True)
        doc.save(str(path))  # type: ignore[attr-defined]
        # Deterministic mtimes, increasing in list order, so v2 FINAL is newer than v1.
        ts = 1_780_000_000 + i * 3600
        os.utime(path, (ts, ts))
        written.append(path)
    return written
