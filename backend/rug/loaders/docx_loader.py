"""Word .docx loader.

Walks the body XML in reading order so paragraphs, tables and images stay interleaved.
Tracked changes are read as accepted: inserted runs (w:ins) are ordinary w:t text and are
kept; deleted runs live in w:delText (or under w:del / w:moveFrom) and are dropped.
Comments, headers and footers live in other parts and are never visited.
"""

import logging
import re
import zipfile
from collections.abc import Iterator
from pathlib import Path

import docx
from docx.document import Document as DocxDocument
from lxml import etree

from rug.config import get_settings
from rug.loaders.base import LoadedDocument, Section
from rug.loaders.ocr import ImageSkipped, ocr_image

log = logging.getLogger(__name__)

W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
A = "http://schemas.openxmlformats.org/drawingml/2006/main"
R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
MC = "http://schemas.openxmlformats.org/markup-compatibility/2006"


def _q(ns: str, tag: str) -> str:
    return f"{{{ns}}}{tag}"


T_TEXT, T_TAB, T_BR, T_CR = _q(W, "t"), _q(W, "tab"), _q(W, "br"), _q(W, "cr")
T_P, T_TBL, T_TR, T_TC, T_SDT = _q(W, "p"), _q(W, "tbl"), _q(W, "tr"), _q(W, "tc"), _q(W, "sdt")
T_BLIP = _q(A, "blip")
# Content under these is not part of the accepted document text.
_EXCLUDED_ANCESTORS = {_q(W, "del"), _q(W, "moveFrom"), _q(MC, "Fallback")}

_HEADING_RE = re.compile(r"^heading\s*(\d)$", re.IGNORECASE)


def _excluded(el: etree._Element, stop: etree._Element) -> bool:
    for anc in el.iterancestors():
        if anc is stop:
            return False
        if anc.tag in _EXCLUDED_ANCESTORS:
            return True
    return False


def _norm(s: str) -> str:
    return re.sub(r"\s+", " ", s).strip()


def _segments(el: etree._Element) -> Iterator[tuple[str, str]]:
    """Yield ("text", chunk) and ("image", rId) in reading order, skipping deleted runs."""
    for node in el.iter(T_TEXT, T_TAB, T_BR, T_CR, T_BLIP):
        if _excluded(node, el):
            continue
        if node.tag == T_BLIP:
            rid = node.get(_q(R, "embed"))
            if rid:
                yield "image", rid
        elif node.tag == T_TEXT:
            yield "text", node.text or ""
        else:
            yield "text", " "


def element_text(el: etree._Element) -> str:
    return _norm("".join(v for kind, v in _segments(el) if kind == "text"))


def _image_rids(el: etree._Element) -> list[str]:
    return [v for kind, v in _segments(el) if kind == "image"]


class DocxLoader:
    extensions: tuple[str, ...] = (".docx",)

    def looks_valid(self, head: bytes) -> bool:
        return head.startswith(b"PK\x03\x04")

    def validate(self, path: Path) -> None:
        if not zipfile.is_zipfile(path):
            raise ValueError("not a valid .docx (zip) file")
        budget = get_settings().max_uncompressed_mb * 1024 * 1024
        try:
            with zipfile.ZipFile(path) as z:
                # zipfile never inflates a member past its declared size, so the header sum
                # is a real bound on what parsing can decompress.
                names = set(z.namelist())
                total = sum(info.file_size for info in z.infolist())
        except zipfile.BadZipFile as e:
            raise ValueError("not a valid .docx (zip) file") from e
        if "word/document.xml" not in names:
            raise ValueError("not a Word document (word/document.xml is missing)")
        if total > budget:
            raise ValueError(f"uncompressed size {total >> 20} MB exceeds budget {budget >> 20} MB")

    def load(self, path: Path) -> LoadedDocument:
        self.validate(path)
        return _Walker(docx.Document(str(path))).run()


class _Walker:
    def __init__(self, doc: DocxDocument):
        self.doc = doc
        self.s = get_settings()
        self.style_names = {s.style_id: s.name or "" for s in doc.styles}
        self.headings: list[tuple[int, str]] = []
        self.sections: list[Section] = []
        self.warnings: list[str] = []
        self.style_title: str | None = None
        self.seen_images: set[str] = set()

    @property
    def heading_path(self) -> tuple[str, ...]:
        return tuple(text for _, text in self.headings)

    def run(self) -> LoadedDocument:
        self._walk(self.doc.element.body)
        core = (self.doc.core_properties.title or "").strip()
        if core:
            return LoadedDocument(core, "core", self.sections, self.warnings)
        if self.style_title:
            return LoadedDocument(self.style_title, "style", self.sections, self.warnings)
        return LoadedDocument(None, "filename", self.sections, self.warnings)

    def _walk(self, container: etree._Element) -> None:
        for child in container:
            if child.tag == T_P:
                self._paragraph(child)
            elif child.tag == T_TBL:
                self._table(child)
            elif child.tag == T_SDT:  # content controls wrap ordinary body content
                content = child.find(_q(W, "sdtContent"))
                if content is not None:
                    self._walk(content)

    def _style_name(self, p: etree._Element) -> str:
        style = p.find(f"{_q(W, 'pPr')}/{_q(W, 'pStyle')}")
        if style is None:
            return ""
        return self.style_names.get(style.get(_q(W, "val")) or "", "")

    def _paragraph(self, p: etree._Element) -> None:
        style = self._style_name(p)
        heading = _HEADING_RE.match(style)
        if style.lower() == "title" or heading:
            text = element_text(p)
            if text and heading:
                level = int(heading.group(1))
                while self.headings and self.headings[-1][0] >= level:
                    self.headings.pop()
                self.headings.append((level, text))
            elif text and self.style_title is None:
                self.style_title = text
            elif text:  # later Title-styled paragraphs are ordinary content
                self.sections.append(Section(text, self.heading_path, "text"))
            self._ocr_all(_image_rids(p))
            return
        # Body paragraph: emit text and OCR sections in reading order.
        buf: list[str] = []
        for kind, value in _segments(p):
            if kind == "text":
                buf.append(value)
                continue
            self._flush_text(buf)
            self._ocr(value)
        self._flush_text(buf)

    def _flush_text(self, buf: list[str]) -> None:
        text = _norm("".join(buf))
        buf.clear()
        if text:
            self.sections.append(Section(text, self.heading_path, "text"))

    def _table(self, tbl: etree._Element) -> None:
        # A row containing images closes the current table section so the OCR text lands
        # after that row; the next section repeats the header row.
        header: list[str] | None = None
        body: list[list[str]] = []
        emitted = False

        def flush() -> None:
            nonlocal body, emitted
            if header is not None and (body or not emitted):
                text = _render_table([header, *body])
                self.sections.append(Section(text, self.heading_path, "table"))
                emitted = True
            body = []

        carry: dict[int, str] = {}  # column -> text, for vertically merged cells
        for tr in tbl.iter(T_TR):
            if tr.getparent() is not tbl:  # nested tables are covered by the cell text
                continue
            row: list[str] = []
            col = 0
            for tc in tr.findall(T_TC):
                tc_pr = tc.find(_q(W, "tcPr"))
                span, vmerge = 1, None
                if tc_pr is not None:
                    gs = tc_pr.find(_q(W, "gridSpan"))
                    if gs is not None:
                        span = int(gs.get(_q(W, "val")) or 1)
                    vm = tc_pr.find(_q(W, "vMerge"))
                    if vm is not None:
                        vmerge = vm.get(_q(W, "val")) or "continue"
                text = element_text(tc)
                if vmerge == "continue":
                    text = carry.get(col, "")
                carry[col] = text
                row.append(text)
                col += span
            if any(row):
                if header is None:
                    header = row
                else:
                    body.append(row)
            rids = _image_rids(tr)
            if rids:
                flush()
                self._ocr_all(rids)
        flush()

    def _ocr_all(self, rids: list[str]) -> None:
        for rid in rids:
            self._ocr(rid)

    def _ocr(self, rid: str) -> None:
        part = self.doc.part.related_parts.get(rid)
        if part is None:
            return
        key = str(part.partname)
        if key in self.seen_images:
            return
        if len(self.seen_images) >= self.s.max_images_per_doc:
            if len(self.seen_images) == self.s.max_images_per_doc:
                self.warnings.append(f"only the first {self.s.max_images_per_doc} images OCR'd")
                self.seen_images.add("<limit>")
            return
        self.seen_images.add(key)
        try:
            text = ocr_image(
                part.blob, self.s.ocr_min_px, self.s.ocr_max_pixels, self.s.ocr_timeout_s
            )
        except ImageSkipped as e:
            log.debug("skip image %s: %s", key, e)
            return
        # LoaderEnvironmentError (Tesseract missing) propagates: the indexer retries the file.
        if text:
            self.sections.append(Section(f"[image text] {text}", self.heading_path, "image_text"))


def _render_table(rows: list[list[str]]) -> str:
    """Header row first, then each data row as 'Header: value; ...' so a row chunk keeps
    its column meaning even when split away from the header."""
    header = rows[0]
    lines = [" | ".join(header)]
    for row in rows[1:]:
        if len(row) == len(header) and all(header):
            lines.append("; ".join(f"{h}: {v}" for h, v in zip(header, row, strict=True) if v))
        else:
            lines.append(" | ".join(row))
    return "\n".join(lines)
